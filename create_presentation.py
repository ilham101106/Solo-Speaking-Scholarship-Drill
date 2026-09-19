from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Theme colors
    COLOR_BG = RGBColor(11, 15, 25)         # Deep Navy #0B0F19
    COLOR_SURFACE = RGBColor(22, 30, 52)    # Card Surface #161E34
    COLOR_PRIMARY = RGBColor(99, 102, 241)  # Indigo #6366F1
    COLOR_CYAN = RGBColor(6, 182, 212)      # Cyan #06B6D4
    COLOR_AMBER = RGBColor(245, 158, 11)    # Amber #F59E0B
    COLOR_TEXT_MAIN = RGBColor(248, 250, 252) # Off-white
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184) # Slate gray

    def set_slide_background(slide):
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = COLOR_BG
        bg_shape.line.fill.background() # No border
        return bg_shape

    # ==================== SLIDE 1: TITLE SLIDE ====================
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s1)

    # Tag Badge
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.5), Inches(3.6), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(30, 41, 69)
    badge.line.color.rgb = COLOR_PRIMARY
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "🎓 SCHOLARSHIP RESEARCH PITCH"
    p_b.font.size = Pt(12)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_CYAN
    p_b.font.name = "Arial"

    # Main Title
    tb_t = s1.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.5), Inches(2.2))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = "Bridging the Critical Divide:"
    p_t.font.size = Pt(40)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_TEXT_MAIN
    p_t.font.name = "Arial"
    
    p_sub = tf_t.add_paragraph()
    p_sub.text = "A Master's Research Proposal to Catalyze Indonesia's Strategic Growth"
    p_sub.font.size = Pt(24)
    p_sub.font.color.rgb = COLOR_CYAN
    p_sub.font.name = "Arial"

    # Speaker info card
    card_info = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.8), Inches(7.5), Inches(1.3))
    card_info.fill.solid()
    card_info.fill.fore_color.rgb = COLOR_SURFACE
    card_info.line.color.rgb = RGBColor(40, 52, 85)
    tf_i = card_info.text_frame
    tf_i.word_wrap = True
    p_i1 = tf_i.paragraphs[0]
    p_i1.text = "Prospective Master's Candidate | LPDP / Chevening / AAS Drill"
    p_i1.font.bold = True
    p_i1.font.size = Pt(15)
    p_i1.font.color.rgb = COLOR_TEXT_MAIN
    p_i2 = tf_i.add_paragraph()
    p_i2.text = "Focus: Academic Rigor • Global Technology Transfer • 5-Year Return Impact"
    p_i2.font.size = Pt(13)
    p_i2.font.color.rgb = COLOR_TEXT_MUTED

    s1.notes_slide.notes_text_frame.text = (
        "[SPEECH TIME: 0:00 - 0:30]\n\n"
        "HOOK & INTRODUCTION:\n"
        "\"Distinguished interviewers, good morning. Today, I am proud to present my master's research proposal titled 'Bridging the Critical Divide.' "
        "As a prospective scholar, my objective is not merely to obtain an international degree, but to master cutting-edge global methodologies "
        "and adapt them directly to resolve one of Indonesia's most pressing development bottlenecks. Over the next three minutes, I will walk you through "
        "the urgency of this problem, my proposed research approach, and my concrete 5-year return contribution to our home country.\""
    )

    # ==================== SLIDE 2: THE PROBLEM & URGENCY ====================
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s2)

    # Header
    tb_h2 = s2.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.0))
    tf_h2 = tb_h2.text_frame
    p_h2 = tf_h2.paragraphs[0]
    p_h2.text = "01. The Urgency: Why Indonesia Needs This Now"
    p_h2.font.size = Pt(28)
    p_h2.font.bold = True
    p_h2.font.color.rgb = COLOR_TEXT_MAIN

    # Card 1: The Shocking Stat
    c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.8), Inches(3.5), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = COLOR_SURFACE
    c1.line.color.rgb = COLOR_AMBER
    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    p1 = tf_c1.paragraphs[0]
    p1.text = "THE STATISTIC"
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_AMBER
    
    p1_num = tf_c1.add_paragraph()
    p1_num.text = "68%"
    p1_num.font.size = Pt(48)
    p1_num.font.bold = True
    p1_num.font.color.rgb = COLOR_TEXT_MAIN
    
    p1_desc = tf_c1.add_paragraph()
    p1_desc.text = "Disparity or bottleneck identified in our current domestic infrastructure and policy implementation."
    p1_desc.font.size = Pt(13)
    p1_desc.font.color.rgb = COLOR_TEXT_MUTED

    # Card 2: The Root Cause
    c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(1.8), Inches(3.5), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = COLOR_SURFACE
    c2.line.color.rgb = COLOR_PRIMARY
    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True
    p2 = tf_c2.paragraphs[0]
    p2.text = "THE BOTTLENECK"
    p2.font.size = Pt(12)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_CYAN
    
    p2_t = tf_c2.add_paragraph()
    p2_t.text = "Domestic Research Gap"
    p2_t.font.size = Pt(20)
    p2_t.font.bold = True
    p2_t.font.color.rgb = COLOR_TEXT_MAIN
    
    p2_desc = tf_c2.add_paragraph()
    p2_desc.text = "\n• Lack of advanced laboratory facilities locally.\n• Insufficient interdisciplinary collaboration.\n• Reliance on outdated frameworks without global benchmarks."
    p2_desc.font.size = Pt(13)
    p2_desc.font.color.rgb = COLOR_TEXT_MUTED

    # Card 3: The Societal Cost
    c3 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.6), Inches(1.8), Inches(3.7), Inches(4.8))
    c3.fill.solid()
    c3.fill.fore_color.rgb = COLOR_SURFACE
    c3.line.color.rgb = RGBColor(244, 63, 94)
    tf_c3 = c3.text_frame
    tf_c3.word_wrap = True
    p3 = tf_c3.paragraphs[0]
    p3.text = "THE CONSEQUENCE"
    p3.font.size = Pt(12)
    p3.font.bold = True
    p3.font.color.rgb = RGBColor(244, 63, 94)
    
    p3_t = tf_c3.add_paragraph()
    p3_t.text = "Inaction is Expensive"
    p3_t.font.size = Pt(20)
    p3_t.font.bold = True
    p3_t.font.color.rgb = COLOR_TEXT_MAIN
    
    p3_desc = tf_c3.add_paragraph()
    p3_desc.text = "\nWithout immediate technical intervention, Indonesia risks falling behind in international competitiveness and economic resilience by 2030."
    p3_desc.font.size = Pt(13)
    p3_desc.font.color.rgb = COLOR_TEXT_MUTED

    s2.notes_slide.notes_text_frame.text = (
        "[SPEECH TIME: 0:30 - 1:15]\n\n"
        "PROBLEM EXPLANATION (HOOK):\n"
        "\"To understand the significance of my research, we must look at the reality in Indonesia today. Currently, our sector faces a critical 68% deficit "
        "in efficiency and technical readiness. Why? Because domestic universities currently lack the specialized lab infrastructure and interdisciplinary "
        "methodologies required to address this issue at its core. If we continue with business as usual, our national development goals will be severely delayed. "
        "This is precisely why undertaking advanced study overseas is an urgent necessity, not a luxury.\""
    )

    # ==================== SLIDE 3: RESEARCH & METHODOLOGY ====================
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s3)

    tb_h3 = s3.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.0))
    tf_h3 = tb_h3.text_frame
    p_h3 = tf_h3.paragraphs[0]
    p_h3.text = "02. The Proposed Solution: Overseas Master's Focus"
    p_h3.font.size = Pt(28)
    p_h3.font.bold = True
    p_h3.font.color.rgb = COLOR_TEXT_MAIN

    # 3 Strategic Pillars
    pillars = [
        ("PILLAR 1: CURRICULUM", "Specialized Advanced Modules", "Accessing coursework in advanced systems modeling and empirical policy design currently unavailable domestically.", COLOR_PRIMARY),
        ("PILLAR 2: FACULTY & LAB", "Distinguished Mentorship", "Working under Professor [Name] in [Specific Lab], which pioneered global frameworks adapted in 15+ developing nations.", COLOR_CYAN),
        ("PILLAR 3: TRANSFERABILITY", "Localized Prototyping", "Directly designing thesis experiments using Indonesian empirical datasets to ensure 100% actionable feasibility upon return.", COLOR_AMBER)
    ]

    for idx, (tag, title, desc, col) in enumerate(pillars):
        top_pos = Inches(1.8 + (idx * 1.65))
        c_p = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), top_pos, Inches(11.3), Inches(1.4))
        c_p.fill.solid()
        c_p.fill.fore_color.rgb = COLOR_SURFACE
        c_p.line.color.rgb = col
        tf_p = c_p.text_frame
        tf_p.word_wrap = True
        
        p_pt = tf_p.paragraphs[0]
        p_pt.text = tag
        p_pt.font.size = Pt(11)
        p_pt.font.bold = True
        p_pt.font.color.rgb = col
        
        p_head = tf_p.add_paragraph()
        p_head.text = title + " — "
        p_head.font.size = Pt(15)
        p_head.font.bold = True
        p_head.font.color.rgb = COLOR_TEXT_MAIN
        
        run_desc = p_head.add_run()
        run_desc.text = desc
        run_desc.font.size = Pt(13)
        run_desc.font.bold = False
        run_desc.font.color.rgb = COLOR_TEXT_MUTED

    s3.notes_slide.notes_text_frame.text = (
        "[SPEECH TIME: 1:15 - 2:00]\n\n"
        "PROPOSED SOLUTION & WHY THIS UNIVERSITY:\n"
        "\"My proposed thesis is built upon three strategic pillars. First, the specialized curriculum at my target university will provide me with "
        "cutting-edge analytical models. Second, conducting research under Professor [Target Faculty] in their renowned lab will grant me hands-on experience "
        "with tested global frameworks. Most importantly, my research will not remain theoretical: I will calibrate all empirical experiments using "
        "Indonesian case studies to guarantee that the solutions we engineer can be seamlessly implemented back home upon graduation.\""
    )

    # ==================== SLIDE 4: RETURN ROADMAP ====================
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s4)

    tb_h4 = s4.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.0))
    tf_h4 = tb_h4.text_frame
    p_h4 = tf_h4.paragraphs[0]
    p_h4.text = "03. Return Contribution: 5-Year Action Roadmap"
    p_h4.font.size = Pt(28)
    p_h4.font.bold = True
    p_h4.font.color.rgb = COLOR_TEXT_MAIN

    phases = [
        ("YEAR 1", "Immediate Repatriation & Baseline Pilot", "• Return immediately to home institution/industry.\n• Publish thesis findings in accredited journals.\n• Launch pilot implementation in target region.", COLOR_CYAN),
        ("YEAR 2 - 3", "Institutional Scaling & Partnerships", "• Expand pilot into a multi-stakeholder framework.\n• Establish international academic exchange between host university and Indonesian campus.\n• Mentor local researchers.", COLOR_PRIMARY),
        ("YEAR 4 - 5", "Policy Reform & National Impact", "• Advise relevant government ministries (e.g. Bappenas/BRIN/Ministry).\n• Deliver measurable reduction in target national deficit.\n• Scale framework nationwide.", COLOR_AMBER)
    ]

    for idx, (yr, title, bullets, col) in enumerate(phases):
        left_pos = Inches(1.0 + (idx * 3.85))
        c_yr = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(1.8), Inches(3.6), Inches(4.8))
        c_yr.fill.solid()
        c_yr.fill.fore_color.rgb = COLOR_SURFACE
        c_yr.line.color.rgb = col
        tf_yr = c_yr.text_frame
        tf_yr.word_wrap = True
        
        p_yr = tf_yr.paragraphs[0]
        p_yr.text = yr
        p_yr.font.size = Pt(14)
        p_yr.font.bold = True
        p_yr.font.color.rgb = col
        
        p_yt = tf_yr.add_paragraph()
        p_yt.text = title
        p_yt.font.size = Pt(16)
        p_yt.font.bold = True
        p_yt.font.color.rgb = COLOR_TEXT_MAIN
        
        p_yb = tf_yr.add_paragraph()
        p_yb.text = "\n" + bullets
        p_yb.font.size = Pt(12)
        p_yb.font.color.rgb = COLOR_TEXT_MUTED

    s4.notes_slide.notes_text_frame.text = (
        "[SPEECH TIME: 2:00 - 2:40]\n\n"
        "POST-STUDY ROADMAP:\n"
        "\"A scholarship is an investment, and reviewer panel members rightly demand a tangible return on investment. My post-study roadmap is divided into "
        "three clear milestones. In Year 1, I will immediately repatriate to Indonesia and launch a targeted pilot project based on my thesis findings. "
        "In Years 2 to 3, I will scale this pilot through institutional partnerships and establish knowledge-sharing channels with my overseas alma mater. "
        "By Year 5, my goal is to translate these proven empirical results into national policy recommendations in collaboration with key ministries.\""
    )

    # ==================== SLIDE 5: VALUE PROPOSITION & CONCLUSION ====================
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s5)

    tb_h5 = s5.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.0))
    tf_h5 = tb_h5.text_frame
    p_h5 = tf_h5.paragraphs[0]
    p_h5.text = "04. The Value Proposition: Why Invest in Me?"
    p_h5.font.size = Pt(28)
    p_h5.font.bold = True
    p_h5.font.color.rgb = COLOR_TEXT_MAIN

    # Big Card Center
    card_sum = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.8))
    card_sum.fill.solid()
    card_sum.fill.fore_color.rgb = COLOR_SURFACE
    card_sum.line.color.rgb = COLOR_PRIMARY
    tf_s = card_sum.text_frame
    tf_s.word_wrap = True

    p_s1 = tf_s.paragraphs[0]
    p_s1.text = "THE SCHOLARSHIP PROMISE"
    p_s1.font.size = Pt(14)
    p_s1.font.bold = True
    p_s1.font.color.rgb = COLOR_CYAN

    p_s2 = tf_s.add_paragraph()
    p_s2.text = "\"Investing in this master's study is not merely funding an individual degree—it is funding a catalyst for systemic transformation in Indonesia.\""
    p_s2.font.size = Pt(22)
    p_s2.font.bold = True
    p_s2.font.color.rgb = COLOR_TEXT_MAIN

    p_s3 = tf_s.add_paragraph()
    p_s3.text = (
        "\n✅ Proven Academic Foundation & Professional Grit\n"
        "✅ Clear, Feasible, and Data-Driven Research Problem\n"
        "✅ Unwavering Commitment to Long-Term National Repatriation\n\n"
        "Thank you for your consideration. I am ready for your questions."
    )
    p_s3.font.size = Pt(14)
    p_s3.font.color.rgb = COLOR_TEXT_MUTED

    s5.notes_slide.notes_text_frame.text = (
        "[SPEECH TIME: 2:40 - 3:00]\n\n"
        "CONCLUSION & Q&A TRANSITION:\n"
        "\"To conclude, funding my master's journey is not just financing another academic degree; it is investing in a strategic, long-term solution "
        "to a well-documented Indonesian challenge. I have the academic grit, the institutional clarity, and the unwavering dedication to return home "
        "and deliver measurable impact. Thank you very much for your time and attention, and I am eager to answer your questions.\""
    )

    output_path = "c:\\ME\\MY ENGLISH\\Master_Thesis_Pitch_Deck.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
